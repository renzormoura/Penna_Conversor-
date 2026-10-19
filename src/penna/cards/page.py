"""Desenho de uma página da ficha a partir do modelo."""
from __future__ import annotations

from typing import Dict, Optional, Set, Tuple

from openpyxl.utils import get_column_letter

from ._reportlab import HexColor, black
from .constants import BORDA_ESP, INSET
from .pdf_text import celula_padrao, fonte, medir, preparar_texto, quebrar
from .template import Modelo

def desenhar_pagina(c, modelo: Modelo, sobrescritas: Dict[str, str],
                     r1: int, r2: int,
                     ocultar: Set[int] = frozenset(),
                     suprimir: Set[Tuple[int, int]] = frozenset(),
                     escala: Optional[float] = None) -> float:
    """Desenha as linhas ``r1..r2`` do modelo e devolve o ``y`` (unidades do
    modelo) onde o conteúdo terminou — o bloco de extras começa dali.

    ``ocultar`` são linhas que o paciente não usa (listas fixas do
    formulário): elas saem da ficha e as linhas seguintes sobem, encolhendo
    a página. ``suprimir`` são células que não devem ser desenhadas (a
    metade sem marcação de uma linha de CID que fica na ficha). ``escala``
    reduz o formulário até ele caber numa página A4 (ver escala_da_ficha).
    """
    larg_pag, alt_pag = modelo.pagina
    if escala is None:
        escala = modelo.escala
    xs = modelo.xs
    nc = modelo.colunas_para(sobrescritas)

    c.saveState()
    util_w = larg_pag - modelo.margem_esq - modelo.margem_dir
    desloc = max(util_w - xs[nc] * escala, 0.0) / 2.0
    c.translate(modelo.margem_esq + desloc, 0)
    c.scale(escala, escala)
    topo = (alt_pag - modelo.margem_topo) / escala

    # posições compactadas: linha oculta não ocupa altura, então cada linha
    # visível sobe o que as ocultas acima dela ocupavam.
    visiveis = {r for r in range(r1, r2 + 1)
                if modelo.alturas[r - 1] > 0 and r not in ocultar}
    pos: Dict[int, Tuple[float, float]] = {}
    y_atual = topo
    for linha in range(r1, r2 + 1):
        if linha in visiveis:
            pos[linha] = (y_atual, y_atual - modelo.alturas[linha - 1])
            y_atual = pos[linha][1]
        else:
            pos[linha] = (y_atual, y_atual)

    def y_topo(linha: int) -> float:
        return pos[linha][0]

    def y_base(linha: int) -> float:
        return pos[linha][1]

    def faixa_y(faixa: Tuple[int, int, int, int]):
        """(x0, x1, y1, y0) de uma mesclagem, pulando as linhas ocultas."""
        vis = [r for r in range(max(faixa[0], r1), min(faixa[2], r2) + 1)
               if r in visiveis]
        if not vis:
            return None
        return xs[faixa[1] - 1], xs[faixa[3]], y_topo(vis[0]), y_base(vis[-1])

    # 1) só as bordas das células do modelo (o Excel imprime as bordas, não as
    # linhas de grade da planilha — grade completa "sujava" a ficha)

    # 2/3) fundos e bordas
    padrao = celula_padrao()
    desenhadas: Set[Tuple[int, int]] = set()
    for linha in range(r1, r2 + 1):
        if linha not in visiveis:
            continue
        for coluna in range(1, nc + 1):
            if (linha, coluna) in suprimir:
                continue
            info = modelo.celulas.get((linha, coluna))
            if info is None and f"{get_column_letter(coluna)}{linha}" in sobrescritas:
                info = padrao      # alvo de sobrescrita sem estilo no template
            faixa = modelo.merges.get((linha, coluna))
            if faixa:
                if (linha, coluna) != (faixa[0], faixa[1]):
                    continue
                caixa = faixa_y(faixa)
                if caixa is None:
                    continue
                x0, x1, y1, y0 = caixa
            else:
                x0, x1 = xs[coluna - 1], xs[coluna]
                y1, y0 = y_topo(linha), y_base(linha)
            if info is None:
                continue
            fundo = info["fill"]
            if modelo.condicionais:
                ref = f"{get_column_letter(coluna)}{linha}"
                realce = modelo.realce_condicional(
                    linha, coluna, sobrescritas.get(ref, info["valor"]))
                if realce:
                    fundo = realce
            if fundo:
                c.setFillColor(HexColor(fundo))
                c.rect(x0, y0, x1 - x0, y1 - y0, stroke=0, fill=1)
            for lado, (estilo, cor) in info["bordas"].items():
                esp = BORDA_ESP.get(estilo, 0.6)
                c.setStrokeColor(HexColor(cor))
                c.setLineWidth(esp)
                if lado == "left":
                    c.line(x0, y0, x0, y1)
                elif lado == "right":
                    c.line(x1, y0, x1, y1)
                elif lado == "top":
                    c.line(x0, y1, x1, y1)
                else:
                    c.line(x0, y0, x1, y0)
            desenhadas.add((linha, coluna))

    # 4) textos
    for linha in range(r1, r2 + 1):
        if linha not in visiveis:
            continue
        for coluna in range(1, nc + 1):
            if (linha, coluna) not in desenhadas:
                continue
            info = modelo.celulas.get((linha, coluna)) or padrao
            texto = sobrescritas.get(f"{get_column_letter(coluna)}{linha}",
                                     info["valor"])
            if not texto:
                continue
            faixa = modelo.merges.get((linha, coluna))
            if faixa:
                caixa = faixa_y(faixa)
                if caixa is None:
                    continue
                x0, x1, y1, y0 = caixa
            else:
                x0, x1 = xs[coluna - 1], xs[coluna]
                y1, y0 = y_topo(linha), y_base(linha)
            nome, tam, negrito, italico = info["fonte"]
            nome_pdf = fonte(nome, negrito, italico)
            texto = preparar_texto(texto, nome_pdf)
            al_h, al_v, wrap = info["al"]
            cor_txt = info["cor_fonte"]
            larg = x1 - x0
            alt = y1 - y0

            # texto transbordando para células vazias à direita (Excel) —
            # vizinha vazia no template E sem valor escrito por cima: o
            # Excel corta o texto na borda da célula que recebeu dado
            # (ex.: rótulo QUANTIDADE: transborda para a caixa da quantidade)
            if not wrap and not faixa:
                fim = coluna + 1
                bloqueada = False
                while fim < nc:
                    if (linha, fim) in modelo.merges:
                        bloqueada = True
                        break
                    vizinha = modelo.celulas.get((linha, fim))
                    if vizinha is not None and vizinha["valor"]:
                        bloqueada = True
                        break
                    if sobrescritas.get(
                            f"{get_column_letter(fim)}{linha}"):
                        bloqueada = True
                        break
                    fim += 1
                # parou em célula ocupada: corta na borda esquerda dela;
                # sem bloqueio, segue até o fim da folha (como no Excel)
                larg = (xs[fim - 1] if bloqueada else xs[fim]) - x0

            espaco = max(larg - 2 * INSET, 4)
            linhas_txt = (quebrar(texto, nome_pdf, tam, espaco) if wrap
                          else [texto])
            if not wrap and medir(texto, nome_pdf, tam) > espaco:
                linhas_txt = [texto]          # cortado pelo clip abaixo

            ascent = tam * 0.78
            descent = tam * 0.22
            leading = tam * 1.22
            altura_txt = leading * len(linhas_txt)
            if al_v == "top":
                primeira = y1 - INSET - ascent
            elif al_v == "center":
                primeira = y1 - (alt - altura_txt) / 2 - ascent
            else:  # bottom (padrão do Excel)
                ultima = y0 + INSET + descent
                primeira = ultima - (len(linhas_txt) - 1) * leading

            # área de corte: a célula + vizinhas vazias (transbordamento)
            x_corte = x0 + larg
            c.saveState()
            caminho = c.beginPath()
            caminho.rect(x0 - 0.5, y0 - 0.5, (x_corte - x0) + 1,
                         (y1 - y0) + 1)
            c.clipPath(caminho, stroke=0, fill=0)
            c.setFillColor(HexColor(cor_txt))
            c.setFont(nome_pdf, tam)
            for i, linha_txt in enumerate(linhas_txt):
                base = primeira + i * leading
                if base < y0 - 1:
                    break
                if al_h == "center":
                    x_txt = x0 + (larg - medir(linha_txt, nome_pdf, tam)) / 2
                elif al_h == "right":
                    x_txt = x1 - INSET - medir(linha_txt, nome_pdf, tam)
                else:
                    x_txt = x0 + INSET
                c.drawString(x_txt, base, linha_txt)
            c.restoreState()

    c.restoreState()
    # fim do conteúdo visível: onde o bloco de extras (se houver) começa
    return y_base(max(visiveis)) if visiveis else topo


def cabecalho_rodape(c, modelo: Modelo, numero: int) -> None:
    larg_pag, alt_pag = modelo.pagina
    c.saveState()
    # cabeçalho central do modelo (topo) — Calibri 9, como no Excel
    if modelo.cabecalho_centro:
        nome_pdf = fonte("Calibri", False, False)
        c.setFillColor(black)
        c.setFont(nome_pdf, 9)
        linhas = [ln for ln in str(modelo.cabecalho_centro).split("\n") if ln]
        base = alt_pag - modelo.alt_cabecalho - 9 * 1.2
        for ln in linhas:
            ln = preparar_texto(ln, nome_pdf)
            c.drawCentredString(larg_pag / 2, base, ln)
            base -= 9 * 1.2
    # rodapé direito "Página N" (rodapé do modelo)
    rodape = str(modelo.rodape_direita or "").replace("&P", str(numero))
    if rodape:
        nome_pdf = fonte("Calibri", False, False)
        c.setFillColor(black)
        c.setFont(nome_pdf, 9)
        c.drawRightString(larg_pag - modelo.margem_dir,
                          modelo.alt_rodape * 0.55,
                          preparar_texto(rodape, nome_pdf))
    c.restoreState()
