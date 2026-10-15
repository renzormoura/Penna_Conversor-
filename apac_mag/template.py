from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .columns import APAC_COLUMNS, PATIENT_COLUMNS, PROCEDURE_COLUMNS

OUTPUT_NAME = "APAC_MAG.xlsx"


def _all_procedure_columns() -> list[str]:
    columns = list(PROCEDURE_COLUMNS)
    for index in range(1, 10):
        columns.extend(
            [
                f"Sec{index}_Cod_Procedimento",
                f"Sec{index}_Cbo",
                f"Sec{index}_Quantidade",
            ]
        )
    return columns


def _headers() -> list[str]:
    ordered: list[str] = []
    for column in (*APAC_COLUMNS, *PATIENT_COLUMNS, *_all_procedure_columns()):
        if column not in ordered:
            ordered.append(column)
    return ordered


def _style_sheet(sheet, headers: list[str], tab_color: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.tabColor = tab_color
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 14), 28)
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).coordinate}2"


def _add_table(sheet, name: str, headers: list[str]) -> None:
    end = sheet.cell(2, len(headers)).coordinate
    table = Table(displayName=name, ref=f"A1:{end}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def _add_validations(sheet, headers: list[str]) -> None:
    positions = {header: index for index, header in enumerate(headers, 1)}
    sex_column = positions.get("Sexo")
    if sex_column:
        validation = DataValidation(type="list", formula1='"F,M,I"', allow_blank=True)
        validation.error = "Informe F, M ou I."
        validation.errorTitle = "Sexo invalido"
        validation.prompt = "Selecione um valor da lista."
        sheet.add_data_validation(validation)
        letter = sheet.cell(1, sex_column).column_letter
        validation.add(f"{letter}2:{letter}5000")
    for header in ("Cpf_Paciente", "Cns_Paciente", "Numero_Apac", "Cod_Procedimento"):
        column = positions.get(header)
        if not column:
            continue
        letter = sheet.cell(1, column).column_letter
        sheet.conditional_formatting.add(
            f"{letter}2:{letter}5000",
            FormulaRule(formula=[f'AND({letter}2<>"",NOT(ISNUMBER(--SUBSTITUTE(SUBSTITUTE({letter}2,".",""),"-",""))))'], fill=PatternFill("solid", fgColor="FCE4D6")),
        )


def create_main_workbook(path: Path) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    procedure_sheet = workbook.active
    procedure_sheet.title = "procedimento_mes"
    database_sheet = workbook.create_sheet("banco_de_dados")
    references_sheet = workbook.create_sheet("referencias")
    configuration_sheet = workbook.create_sheet("configuracao")

    headers = _headers()
    procedure_sheet.append(headers)
    procedure_sheet.append([""] * len(headers))
    _style_sheet(procedure_sheet, headers, "1F4E78")
    _add_table(procedure_sheet, "ProcedimentoMes", headers)
    _add_validations(procedure_sheet, headers)

    database_headers = list(dict.fromkeys(PATIENT_COLUMNS))
    database_sheet.append(database_headers)
    database_sheet.append([""] * len(database_headers))
    _style_sheet(database_sheet, database_headers, "548235")
    _add_table(database_sheet, "BancoDeDados", database_headers)

    reference_headers = ["tipo", "codigo", "descricao", "competencia", "ativo", "regra"]
    references_sheet.append(reference_headers)
    references_sheet.append([""] * len(reference_headers))
    _style_sheet(references_sheet, reference_headers, "BF9000")
    _add_table(references_sheet, "Referencias", reference_headers)

    configuration_headers = ["chave", "valor", "descricao"]
    configuration_sheet.append(configuration_headers)
    rows = [
        ["competencia", "", "Competencia no formato AAAAMM"],
        ["uf", "", "Codigo da UF"],
        ["cnes_solicitante", "", "CNES do solicitante"],
        ["cnes_executante", "", "CNES do executante"],
        ["codigo_municipio", "", "Codigo IBGE do municipio"],
        ["codigo_orgao_emissor", "", "Codigo do orgao emissor"],
        ["pasta_exportacao", "Exportados", "Pasta para os arquivos gerados"],
    ]
    for row in rows:
        configuration_sheet.append(row)
    _style_sheet(configuration_sheet, configuration_headers, "8064A2")
    _add_table(configuration_sheet, "Configuracao", configuration_headers)

    workbook.save(destination)
    return destination


if __name__ == "__main__":
    create_main_workbook(Path.cwd() / OUTPUT_NAME)
